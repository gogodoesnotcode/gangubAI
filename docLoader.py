import os
import fitz  # PyMuPDF
from pptx import Presentation
from langchain_core.documents import Document

def process_file_to_chunks(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    ppt_name = os.path.basename(file_path)
    chunks = []

    if ext == ".pptx":
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_num = i + 1
            text = "\n".join([shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()])
            if text:
                chunks.append(create_doc(text, ppt_name, "Slide", slide_num))

    elif ext == ".pdf":
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc):
            page_id = page_num + 1
            text = page.get_text().strip()
            if text:
                chunks.append(create_doc(text, ppt_name, "Page", page_id))
    
    return chunks

def create_doc(content, source_name, unit_type, unit_num):
    """Helper to standardize metadata and inject context."""
    metadata = {
        "source": source_name,
        "type": unit_type,
        "number": unit_num
    }
    # Context Injection: Force the LLM to 'see' the source in the text
    enriched_content = f"[{unit_type} {unit_num} from {source_name}]\nContent: {content}"
    return Document(page_content=enriched_content, metadata=metadata)


all_chunks = process_file_to_chunks("D:/Everything/GangubAI/Unit1/UE23CS351B_1ee2e7d3-eaba-11f0-9f05-063baf11569b_20260202121720.pdf")
# all_chunks += process_file_to_chunks("Robotics_Notes.pdf")

print(all_chunks[0])