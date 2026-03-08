from langgraph.graph import StateGraph, START, END
from typing import Annotated, TypedDict
from langchain_core.messages import BaseMessage, HumanMessage, AIMessage, SystemMessage, ToolMessage
from pydantic import BaseModel, Field
from langchain_groq import ChatGroq
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.tools import tool
from langgraph.prebuilt import ToolNode, tools_condition
import os
import fitz  # PyMuPDF
from pptx import Presentation
from langchain_community.vectorstores import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain.chat_models import init_chat_model
from langgraph.checkpoint.sqlite import SqliteSaver, sqlite3
from dotenv import load_dotenv
from langgraph.graph.message import add_messages

load_dotenv()

class ChatState(TypedDict):
    messages: Annotated[list[BaseMessage], add_messages]

class RAGResponse(BaseModel):
    explanation: str = Field(description="Detailed explanation answering the user's question based on the retrieved context")
    citation: str = Field(description="Citation in the format: 'For more reference check <source>, <type> <number>' e.g. 'For more reference check ECC_1_Introduction.pptx, Slide 3'")

# Document Processing Functions
def process_file_to_chunks(file_path, chunk_size=1000, chunk_overlap=200):
    """Parses PPTX and PDF with text splitting for retrieval."""
    ext = os.path.splitext(file_path)[1].lower()
    file_name = os.path.basename(file_path)
    documents = []

    if ext == ".pptx":
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            text = "\n".join([s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text.strip()])
            if text: 
                doc = Document(
                    page_content=text,
                    metadata={"source": file_name, "type": "Slide", "number": i + 1}
                )
                documents.append(doc)

    elif ext == ".pdf":
        doc = fitz.open(file_path)
        for i, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                page_doc = Document(
                    page_content=text,
                    metadata={"source": file_name, "type": "Page", "number": i + 1}
                )
                documents.append(page_doc)
    
    # Use RecursiveCharacterTextSplitter it keeps the larger units(eg: paragraphs together)
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        add_start_index=True,
    )
    
    # Split documents into smaller chunks
    chunks = text_splitter.split_documents(documents)
    
    # Add the metadata inside the page content so that the llm can use it to print
    for chunk in chunks:
        original_content = chunk.page_content
        metadata = chunk.metadata
        chunk.page_content = (
            f"[{metadata['type']} {metadata['number']} from {metadata['source']}]\n"
            f"Content: {original_content}"
        )
    
    return chunks

# Setup Gangubai RAG System
def setup_gangubai_brain(file_paths, persist_directory=".gangubai_db_hf/"):
    """Initialize the RAG system with document indexing."""
    print("📚 Loading and processing documents...")
    # 1. Load and Chunk all documents
    all_chunks = []
    for path in file_paths:
        if os.path.exists(path):
            print(f"  - Processing {path}...")
            chunks = process_file_to_chunks(path)
            all_chunks.extend(chunks)
            print(f"    ✓ Created {len(chunks)} chunks")
        else:
            print(f"  ⚠️  File not found: {path}")
    
    print(f"\n📊 Total chunks created: {len(all_chunks)}")
    
    # 2. Initialize Embeddings (HuggingFace - BAAI/bge-large-en-v1.5)
    print("🔄 Initializing embeddings...")
    embeddings = HuggingFaceEmbeddings(model_name="BAAI/bge-large-en-v1.5")
    
    # 3. Create Vector Store (ChromaDB)
    print("💾 Creating vector database...")
    vectorstore = Chroma.from_documents(
        documents=all_chunks, 
        collection_name="gangubai_collection",
        embedding=embeddings,
        persist_directory=persist_directory
    )
    
    print(f"✅ Vector database created with {len(all_chunks)} documents\n")
    
    return vectorstore

@tool
def retrieve_context(query: str) -> str:
    """Retrieve information from the Gangubai knowledge base to help answer a query.
    
    Use this tool when you need to answer questions about topics in the knowledge base.
    
    Args:
        query: The question or topic to search for in the knowledge base
        
    Returns:
        Relevant information from the knowledge base with source citations
    """
    retrieved_docs = vectorstore.similarity_search(query, k=3)
    serialized = "\n\n".join(
        (f"Source: {doc.metadata}\nContent: {doc.page_content}")
        for doc in retrieved_docs
    )
    return serialized


@tool
def calculator(first_num: float, second_num: float, operation: str) -> str:
    """Perform basic arithmetic operations.
    
    Use this tool when you need to calculate: addition, subtraction, multiplication, or division.
    
    Args:
        first_num: The first number in the calculation
        second_num: The second number in the calculation  
        operation: Must be one of: "add", "subtract", "multiply", "divide"
        
    Returns:
        A string with the calculation result
    """
    try:
        if operation == "add":
            result = first_num + second_num
        elif operation == "subtract":
            result = first_num - second_num
        elif operation == "multiply":
            result = first_num * second_num
        elif operation == "divide":
            if second_num == 0:
                return "Error: Cannot divide by zero"
            result = first_num / second_num
        else:
            return f"Error: Invalid operation '{operation}'. Use: add, subtract, multiply, or divide"
        return f"The result of {first_num} {operation} {second_num} is {result}"
    except Exception as e:
        return f"Calculation error: {str(e)}"


# Initialize or load existing vector store
# Option 1: Create new vectorstore from files
my_files = [
    "Unit1\ECC_1_Introduction to GenAI and Its Application.pptx",
    "Unit1\ECC_2_LLM basics and Evolution.pptx",
    "Unit1\ECC_3_NLP Basics_Terminologies_TaskOverview.pptx"
]

# Option 2: Load existing vectorstore if it exists
persist_directory = ".gangubai_db_hf/"

if os.path.exists(persist_directory):
    print("Loading existing vector database...")
    embeddings = HuggingFaceEmbeddings(model_name="BAAI/bge-large-en-v1.5")
    vectorstore = Chroma(
        collection_name="gangubai_collection",
        embedding_function=embeddings,
        persist_directory=persist_directory
    )
    print("✅ Vector database loaded\n")
else:
    vectorstore = setup_gangubai_brain(my_files, persist_directory)

# Configure tools and LLM
tools = [calculator, retrieve_context]

model = init_chat_model("google_genai:gemini-2.5-flash-lite")

system_prompt = (
        "You are GangubAI, a helpful AI assistant with access to a knowledge base. "
        "Use the retrieve_context tool to search for relevant information before answering questions. "
        "Always cite your sources using the metadata provided (source, type, and number). "
        "If you can't find relevant information, say so honestly."
    )

# Bind tools to the model - LangGraph will handle the agent logic
llm_with_tools = model.bind_tools(tools)

# Structured output LLM used only after RAG retrieval
llm_for_rag = model.with_structured_output(RAGResponse)

def chatNode(state: ChatState) -> ChatState:
    messages = state['messages']
    all_messages = [SystemMessage(content=system_prompt)] + list(messages)

    # Check if retrieve_context was called in the most recent tool round-trip
    rag_was_called = False
    for msg in reversed(messages):
        if isinstance(msg, AIMessage):
            break
        if isinstance(msg, ToolMessage) and msg.name == "retrieve_context":
            rag_was_called = True
            break

    if rag_was_called:
        response: RAGResponse = llm_for_rag.invoke(all_messages)
        formatted = f"{response.explanation}\n\n**For more reference:** {response.citation}"
        return {'messages': [AIMessage(content=formatted)]}

    response = llm_with_tools.invoke(all_messages)
    return {'messages': [response]}

tool_node = ToolNode(tools)

graph = StateGraph(ChatState)

graph.add_node('chat', chatNode)
graph.add_node('tools', tool_node)

graph.add_edge(START, 'chat')

graph.add_conditional_edges('chat', tools_condition)
graph.add_edge('tools', 'chat')  # After tools execute, go back to chat

# For LangGraph API - compile without checkpointer (API handles persistence)
agent = graph.compile()


# Interactive Mode - GangubAI Chatbot with RAG
if __name__ == "__main__":
    # For standalone mode - use local checkpointer
    conn = sqlite3.connect(database='gangubai_chatbot_history.db', check_same_thread=False)
    checkpointer = SqliteSaver(conn=conn)
    chatbot = graph.compile(checkpointer=checkpointer)
    
    print("\n" + "="*60)
    print("💬 GangubAI Chatbot - Type 'quit' or 'exit' to stop")
    print("="*60)

    config = {'configurable': {'thread_id': '2'}}

    while True:
        try:
            user_query = input("\n❓ Your question: ").strip()
            
            if user_query.lower() in ['quit', 'exit', 'q', 'bye']:
                print("\n👋 Goodbye! GangubAI shutting down...")
                break
            
            if not user_query:
                continue
            
            print(f"\n{'='*60}")
            print(f"❓ Question: {user_query}")
            print('='*60)
            print("\n💭 GangubAI is thinking...\n")
            
            # Stream the response
            for event in chatbot.stream(
                {'messages': [HumanMessage(content=user_query)]}, 
                config=config,
                stream_mode="values"
            ):
                event["messages"][-1].pretty_print()
            
            print('='*60)
            
        except KeyboardInterrupt:
            print("\n\n👋 Goodbye! GangubAI shutting down...")
            break
        except Exception as e:
            print(f"\n❌ Error: {e}")
            print("Please try again with a different question.")

